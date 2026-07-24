"""端到端测试：按物料编号关联入库 + 新价格推导 + 多页报价单生成 + 公式校验。"""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import openpyxl
from openpyxl.utils import get_column_letter

import zd_config as C
import zd_db as DB
import zd_excel as XL
from samples import make_samples

PASS = []
FAIL = []


def check(name, cond, extra=""):
    (PASS if cond else FAIL).append(name)
    print(("✅" if cond else "❌"), name, extra)


def main():
    make_samples.make()
    DB.init_db()
    DB.clear_db()

    cfg = os.path.join(C.SAMPLES_DIR, "配单表.xlsx")
    prc = os.path.join(C.SAMPLES_DIR, "价格表.xlsx")
    qte = os.path.join(C.SAMPLES_DIR, "报价表.xlsx")
    enc_prc = os.path.join(C.SAMPLES_DIR, "价格表_加密.xlsx")
    std_enc_prc = os.path.join(C.SAMPLES_DIR, "价格表_标准加密.xlsx")

    # 1) 更新数据库（按物料编号关联 + 批量 upsert）
    res = XL.update_database(cfg, prc)
    check("入库行数=4（含特殊编码）", res["rows"] == 4, f"got {res['rows']}")
    check("自适应线程数≥1", res["workers"] >= 1, f"workers={res['workers']}")
    check("配单表更新日期已记录", res["config_date"] not in (None, "—"))
    check("价格表更新日期已记录", res["price_date"] not in (None, "—"))

    # 1b) 带密码的价格表（敏捷加密）：自动解密后入库
    res_e = XL.update_database(cfg, enc_prc)
    check("加密价格表(敏捷)自动解密入库=4", res_e["rows"] == 4, f"got {res_e['rows']}")

    # 1b') 带密码的价格表（**标准加密 / CFB 容器**，即真实「带密码测试.xlsx」那类）
    res_se = XL.update_database(cfg, std_enc_prc)
    check("加密价格表(标准/CFB)自动解密入库=4", res_se["rows"] == 4, f"got {res_se['rows']}")

    # 1c) 特殊物料编码 1200000043：市场价=15000，其余=9000（强制覆盖）
    M = "市场价"
    check("1200000043 市场价=15000",
          DB.get_price("小米", "SP-999", M) == 15000.0)
    check("1200000043 H价=9000",
          DB.get_price("小米", "SP-999", "核心价（H价）") == 9000.0)
    check("1200000043 D价=9000",
          DB.get_price("小米", "SP-999", "大客户价（D价）") == 9000.0)
    check("1200000043 Y价=9000",
          DB.get_price("小米", "SP-999", "业务底价（Y价）") == 9000.0)
    check("1200000043 批发价=9000",
          DB.get_price("小米", "SP-999", "批发价（P价）") == 9000.0)

    # 2) 价格推导（按物料编号关联；Y 基准、向上取整）
    H = "核心价（H价）"; Q = "签约价（Q价）"; D = "大客户价（D价）"
    P = "批发价（P价）"; Y = "业务底价（Y价）"
    M = "市场价"; B = "概算价"; S = "销售指导价"; L = "租赁价"

    check("FW-100 H=1000(导入)", DB.get_price("华为", "FW-100", H) == 1000.0)
    check("FW-100 Y=ceil(1000/0.63)=1588",
          DB.get_price("华为", "FW-100", Y) == 1588)
    check("FW-100 Q=ceil(1588*0.7)=1112",
          DB.get_price("华为", "FW-100", Q) == 1112)
    check("FW-100 D=ceil(1588*0.8)=1271",
          DB.get_price("华为", "FW-100", D) == 1271)
    check("FW-100 市场=ceil(1588*2)=3176",
          DB.get_price("华为", "FW-100", M) == 3176)
    check("FW-100 P=900(导入,不推算)",
          DB.get_price("华为", "FW-100", P) == 900.0)
    check("FW-100 概算价留空", DB.get_price("华为", "FW-100", B) is None)

    check("SW-200 Q=700(导入)", DB.get_price("华为", "SW-200", Q) == 700)
    check("SW-200 Y=ceil(700/0.7)=1000",
          DB.get_price("华为", "SW-200", Y) == 1000)
    check("SW-200 H=ceil(1000*0.63)=630",
          DB.get_price("华为", "SW-200", H) == 630)
    check("SW-200 D=800", DB.get_price("华为", "SW-200", D) == 800)

    check("SRV-300 市场=5000(导入)", DB.get_price("戴尔", "SRV-300", M) == 5000)
    check("SRV-300 Y=ceil(5000/2)=2500",
          DB.get_price("戴尔", "SRV-300", Y) == 2500)
    check("SRV-300 H=ceil(2500*0.63)=1575",
          DB.get_price("戴尔", "SRV-300", H) == 1575)
    check("SRV-300 D=ceil(2500*0.8)=2000",
          DB.get_price("戴尔", "SRV-300", D) == 2000)
    check("SRV-300 概算价=6000(直接价原样)",
          DB.get_price("戴尔", "SRV-300", B) == 6000.0)
    check("SRV-300 销售指导价=7000(原样)",
          DB.get_price("戴尔", "SRV-300", S) == 7000.0)
    check("SRV-300 租赁价=4000(原样)",
          DB.get_price("戴尔", "SRV-300", L) == 4000.0)

    # 3) 生成多页报价单（大客户价 D，比例 0.9）
    rows = XL.read_workbook(qte)
    out, ver, nf, nm, pages = XL.write_quote(rows, D, 0.9, qte)
    check("产出文件存在", os.path.exists(out), out)
    check("生成 2 个分表页", pages == 2, f"pages={pages}")
    check("跨页命中 4 行 / 未找到 2 行", nf == 4 and nm == 2, f"nf={nf}, nm={nm}")

    # 4) 校验多页结构 + 公式与单价
    wb = openpyxl.load_workbook(out)
    check("含『报价』页", "报价" in wb.sheetnames, str(wb.sheetnames))
    check("含『报价二』页", "报价二" in wb.sheetnames, str(wb.sheetnames))

    ws = wb["报价"]
    cols = {ws.cell(row=1, column=c).value: c
            for c in range(1, ws.max_column + 1)}
    price_c = get_column_letter(cols["单价"])
    he_c = get_column_letter(cols["合价"])

    def rowdict(ws, r):
        return {ws.cell(row=1, column=c).value: ws.cell(row=r, column=c).value
                for c in range(1, ws.max_column + 1)}

    d2 = rowdict(ws, 2)
    check("报价页 FW-100 单价=1143.9", d2["单价"] == 1143.9, f"got {d2['单价']}")
    check("报价页 FW-100 合价为公式", str(d2["合价"]).startswith("=IF("),
          str(d2["合价"]))
    # 标底参数来自配单表（存于 DB），报价表本身不含该列也应带出
    check("报价页 FW-100 标底参数=参数A", d2["标底参数"] == "参数A",
          f"got {d2.get('标底参数')}")
    d5 = rowdict(ws, 5)
    check("报价页 XXXX 单价留空", d5["单价"] in (None, ""), f"got {d5['单价']}")
    last = ws.max_row
    total_f = ws.cell(row=last, column=cols["合价"]).value
    check("报价页 总价=SUM公式", str(total_f).startswith("=SUM("),
          str(total_f))

    ws2 = wb["报价二"]
    cols2 = {ws2.cell(row=1, column=c).value: c
             for c in range(1, ws2.max_column + 1)}
    d2b = rowdict(ws2, 2)
    check("报价二页 FW-100 单价=1143.9", d2b["单价"] == 1143.9, f"got {d2b['单价']}")
    last2 = ws2.max_row
    total2 = ws2.cell(row=last2, column=cols2["合价"]).value
    check("报价二页 各自总价=SUM公式", str(total2).startswith("=SUM("),
          str(total2))
    d3 = rowdict(ws2, 3)
    check("报价二页 XXXX 单价留空", d3["单价"] in (None, ""), f"got {d3['单价']}")

    print("\n--- 结果 ---")
    print(f"通过 {len(PASS)} / 失败 {len(FAIL)}")
    if FAIL:
        print("失败项：", FAIL)
        raise SystemExit(1)
    print("🎉 全部通过！报价单：", out)


if __name__ == "__main__":
    main()
