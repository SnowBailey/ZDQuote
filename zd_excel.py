"""ZD 报价助手 - Excel 读写与清洗。

- 多 sheet 读取，表头智能识别（最长同义词匹配）
- 清洗：跳过空行/无型号行，图片在 openpyxl 读取时天然被忽略
- 品牌缺省时回退为 sheet 名（适配“多品牌页”）
- 并行清洗：每个 sheet 由独立线程打开只读 workbook 解析，I/O 释放 GIL，
  线程数按本机 CPU 核心自适应（适配 M1/M2/M3/M4）
- 写报价单：12 列、合价/总价用 Excel 公式、未找到型号单价留空；
  源表含多页时，逐页对应生成多张输出 sheet
"""
import os
import io
import math
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from concurrent.futures import ThreadPoolExecutor, as_completed

import zd_config as C
import zd_pricing as P
import zd_db as DB
import zd_decrypt as DEC


def scaled(progress, lo, hi):
    """返回一个子进度回调：把内层 [0,1] 映射到外层 [lo,hi]。"""
    if progress is None:
        return None
    def cb(f, t):
        progress(lo + (hi - lo) * f, t)
    return cb


# ---------- 工具 ----------
def _norm(s):
    return "".join(str(s).lower().split())


def _to_number(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("¥", "").replace("￥", "")
    if s == "":
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _map_headers(header_cells):
    """返回 {field: col_index(0-based)}，最长同义词优先。"""
    syn = {f: [_norm(x) for x in lst] for f, lst in C.FIELD_SYNONYMS.items()}
    result = {}
    for ci, cell in enumerate(header_cells):
        if cell is None:
            continue
        h = _norm(cell)
        if not h:
            continue
        best_field = None
        best_len = -1
        for f, lst in syn.items():
            if f in result:
                continue
            for sub in lst:
                if sub and sub in h and len(sub) > best_len:
                    best_len = len(sub)
                    best_field = f
        if best_field:
            result[best_field] = ci
    return result


def _emit_row(row, mp, sheet_name, out):
    """依据表头映射从一行（tuple）抽取字段，追加到 out。

    数据行判定：有「型号」或「物料编号」之一即视为数据行
    （配单表/报价表常有型号；价格表可能仅靠物料编号定位）。
    """
    raw = {}
    for f, ci in mp.items():
        v = row[ci] if ci < len(row) else None
        if v is None or (isinstance(v, str) and v.strip() == ""):
            continue
        raw[f] = v
    if not raw:
        return
    model = (str(raw.get("model", "")).strip()
            if raw.get("model") not in (None, "") else "")
    code = (str(raw.get("code", "")).strip()
           if raw.get("code") not in (None, "") else "")
    if not model and not code:
        return  # 既无型号也无物料编号，视为非数据行
    brand = str(raw.get("brand", "")).strip() or str(sheet_name).strip()
    raw["brand"] = brand
    raw["model"] = model
    raw["code"] = code
    for numf in ("qty", *C.PRICE_RAW_FIELDS):
        if numf in raw:
            nv = _to_number(raw[numf])
            if nv is not None:
                raw[numf] = nv
    raw["_sheet"] = sheet_name
    out.append(raw)


def _clean_sheet(src, sheet_name):
    """打开独立的只读 workbook，用 iter_rows 流式清洗单个 sheet。

    src 可以是文件路径（str）或已解密的内存 BytesIO。
    关键性能点：read_only 模式下 *切勿* 用 ws.cell(row, col) 随机访问，
    那会让 openpyxl 反复重定位、速度呈数量级下降。必须用 iter_rows
    一次性流式读完。表头扫描仍在流内完成（缓存前若干行），不回头 seek。
    """
    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    try:
        ws = wb[sheet_name]
    except KeyError:
        wb.close()
        return []
    nrows = ws.max_row
    scan = min(nrows, 20)
    cache = []
    best_row, best_score = 1, -1
    for ri, row in enumerate(
        ws.iter_rows(min_row=1, max_row=scan, values_only=True), start=1
    ):
        cache.append(row)
        mp = _map_headers(list(row))
        score = len(mp)
        if score > best_score:
            best_score, best_row = score, ri
    if best_score <= 0:
        wb.close()
        return []
    mp = _map_headers(list(cache[best_row - 1]))
    out = []
    # 处理扫描区中、表头之后的数据行（已缓存，不重复读）
    for row in cache[best_row:]:
        _emit_row(row, mp, sheet_name, out)
    # 读取扫描区之后的全部数据行
    if scan < nrows:
        for row in ws.iter_rows(min_row=scan + 1, values_only=True):
            _emit_row(row, mp, sheet_name, out)
    wb.close()
    return out


def read_workbook(path, max_workers=None, progress=None):
    """并行读取所有 sheet，返回行 dict 列表（含字段 + '_sheet'，按 sheet 顺序）。
    progress(frac, text) 在每清洗完一个 sheet 时上报（0~1）。

    若源文件为带密码的 Excel（ECMA-376 Agile 加密），先用默认密码
    PRICE_FILE_PASSWORD 解密到内存再读取；非加密文件原样处理。
    """
    if max_workers is None:
        max_workers = C.cpu_workers()
    # 解密（带密码的 xlsx 返回字节，否则 None）
    data = DEC.decrypt_xlsx(path, C.PRICE_FILE_PASSWORD)
    src = io.BytesIO(data) if data is not None else path
    wb0 = openpyxl.load_workbook(src, read_only=True)
    names = list(wb0.sheetnames)
    wb0.close()
    total = len(names)

    if total <= 1 or max_workers <= 1:
        rows = []
        for i, n in enumerate(names, 1):
            rows.extend(_clean_sheet(src, n))
            if progress:
                progress(i / total, f"清洗表页 {i}/{total}：{n}")
        return rows

    per = {}
    done = 0
    with ThreadPoolExecutor(max_workers=max_workers) as ex:
        futs = {ex.submit(_clean_sheet, src, n): n for n in names}
        for fut in as_completed(futs):
            n = futs[fut]
            per[n] = fut.result()
            done += 1
            if progress:
                progress(done / total, f"清洗表页 {done}/{total}：{n}")
    ordered = []
    for n in names:  # 维持原 sheet 顺序
        ordered.extend(per.get(n, []))
    return ordered


def _norm_code(v):
    """物料编号归一化：去空格、转小写，便于按 code 关联两张表。"""
    if v is None:
        return ""
    return "".join(str(v).lower().split())


def build_merged_rows(config_path, price_path, progress=None):
    """读取配单表 + 价格表，按「物料编号(code)」关联并推导价格。

    规则：
    - 配单表为主（含品牌/型号/物料编号/描述字段）；
    - 价格表按物料编号建索引，逐行关联到配单表对应设备；
    - 命中的价格按 derive_prices 推导（直接价原样、派生价补缺失，向上取整）；
    - 输出报价仍以品牌+型号为基础（DB 主键保持 brand+model）。
    """
    cfg = read_workbook(config_path, progress=scaled(progress, 0.0, 0.35))
    prc = read_workbook(price_path, progress=scaled(progress, 0.35, 0.70))

    if progress:
        progress(0.72, "按物料编号关联配单与价格…")

    # 价格表按物料编号建索引（同 code 仅取首个）
    price_by_code = {}
    for r in prc:
        code = _norm_code(r.get("code"))
        if not code:
            continue
        if code not in price_by_code:
            price_by_code[code] = {k: r.get(k) for k in C.PRICE_RAW_FIELDS}

    # 以配单表为主，逐行关联价格
    merged = {}
    for r in cfg:
        brand = str(r.get("brand", "")).strip() or str(r.get("_sheet", "")).strip()
        model = str(r.get("model", "")).strip()
        if not model:
            continue
        key = (brand, model)
        rec = merged.setdefault(key, {"brand": brand, "model": model})
        for f in ("seq", "name", "spec", "country", "qty", "unit",
                   "remark", "code", "bid_param"):
            if r.get(f) not in (None, ""):
                rec[f] = r[f]
        code = _norm_code(r.get("code"))
        if code and code in price_by_code:
            derived = P.derive_prices(price_by_code[code])
            # 特殊物料编码：强制覆盖价格（无视价格表导入值）
            if code in C.SPECIAL_PRICE_OVERRIDE:
                derived = dict(C.SPECIAL_PRICE_OVERRIDE[code])
            for k, v in derived.items():
                if v is not None:
                    rec[k] = v

    if progress:
        progress(0.85, f"合并完成：{len(merged)} 条设备")
    return list(merged.values())


def update_database(config_path, price_path, progress=None):
    """完整更新流程：并行清洗→推导→批量入库→记录两源文件更新日期。

    每次导入都先清空老库（配单/价格表已换新，旧数据不再适用）。
    """
    DB.init_db()
    DB.clear_db()  # 清空旧设备与源日期，再写入本次数据
    nw = C.cpu_workers()
    if progress:
        progress(0.02, "开始清洗配单表与价格表…")
    rows = build_merged_rows(config_path, price_path,
                             progress=scaled(progress, 0.02, 0.85))
    if progress:
        progress(0.88, f"批量写入数据库（{nw} 线程）…")
    n = DB.bulk_upsert(rows)
    if progress:
        progress(0.97, "记录源文件更新日期…")
    # 记录源文件更新日期（文件修改时间）
    for path, key in ((config_path, "config_updated"),
                      (price_path, "price_updated")):
        mt = os.path.getmtime(path)
        ds = datetime.datetime.fromtimestamp(mt).strftime("%Y-%m-%d %H:%M")
        DB.set_meta(key, ds)
    if progress:
        progress(1.0, "数据库更新完成")
    return {"rows": n, "workers": nw,
            "config_date": DB.get_meta("config_updated"),
            "price_date": DB.get_meta("price_updated")}


# ---------- 写报价单 ----------
_HDR_FILL = PatternFill("solid", fgColor="2E5AAC")
_HDR_FONT = Font(color="FFFFFF", bold=True, size=11)
_TOTAL_FILL = PatternFill("solid", fgColor="DCE6F7")
_THIN = Side(style="thin", color="BBBBBB")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

# Excel 工作表名非法字符
_BAD = str.maketrans({c: "_" for c in ':\\/?*[]'})


def _safe_sheet_title(name, existing: list) -> str:
    """清洗为合法、唯一、≤31 字符的 sheet 名（保留顺序靠调用方）。"""
    if not name or str(name).strip() in ("", "Sheet", "Sheet1"):
        base = "报价单"
    else:
        base = str(name).translate(_BAD).strip()
    base = base[:31] if base else "报价单"
    if base not in existing:
        return base
    i = 2
    while f"{base}_{i}"[:31] in existing:
        i += 1
    return f"{base}_{i}"[:31]


def _write_sheet(ws, rows, fields, labels, price_idx, he_idx, qty_idx,
                 name_idx, system_label, ratio):
    """写入单张报价 sheet，返回 (命中数, 未找到数)。"""
    DB.init_db()
    # 表头
    for ci, lab in enumerate(labels, start=1):
        c = ws.cell(row=1, column=ci, value=lab)
        c.fill = _HDR_FILL
        c.font = _HDR_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = _BORDER
    ws.freeze_panes = "A2"

    n_found = 0
    n_missing = 0
    r = 2
    # 报价表仅含品牌/型号/数量/单位等定价所需字段；
    # 标底参数、规格、国别、备注、序号、设备名称等来自配单表（存于 DB）。
    # 输出时优先用报价表行自身的值，缺省字段回退到 DB 记录对应设备。
    _FALLBACK = ("seq", "name", "spec", "country", "remark", "bid_param")
    for row in rows:
        brand = row.get("brand")
        model = row.get("model")
        price_val = DB.get_price(brand, model, system_label)
        if price_val is not None:
            unit_price = round(price_val * ratio, 2)
            n_found += 1
        else:
            unit_price = None
            n_missing += 1
        rec = DB.get_record(brand, model) if (brand and model) else None
        for ci, f in enumerate(fields, start=1):
            if f == "price":
                val = unit_price
            elif f == "he":
                val = None  # 公式稍后填
            else:
                val = row.get(f)
                if (val is None or val == "") and rec is not None and f in _FALLBACK:
                    val = rec.get(f)
            cell = ws.cell(row=r, column=ci, value=val)
            cell.border = _BORDER
            if f == "qty" and isinstance(val, (int, float)):
                cell.number_format = "0.##"
            if f == "price" and isinstance(val, (int, float)):
                cell.number_format = "0.00"
        # 合价公式：无单价则留空
        price_col = get_column_letter(price_idx + 1)
        qty_col = get_column_letter(qty_idx + 1)
        he_cell = ws.cell(row=r, column=he_idx + 1)
        he_cell.value = f'=IF({price_col}{r}="","",{qty_col}{r}*{price_col}{r})'
        he_cell.number_format = "0.00"
        he_cell.border = _BORDER
        r += 1

    last = r - 1
    # 总价行（本页）
    he_col = get_column_letter(he_idx + 1)
    tcell = ws.cell(row=r, column=name_idx + 1, value="总价")
    tcell.font = Font(bold=True)
    tcell.fill = _TOTAL_FILL
    tcell.border = _BORDER
    total = ws.cell(row=r, column=he_idx + 1,
                    value=f"=SUM({he_col}2:{he_col}{last})")
    total.font = Font(bold=True)
    total.number_format = "0.00"
    total.fill = _TOTAL_FILL
    total.border = _BORDER

    widths = [8, 26, 12, 18, 30, 8, 8, 7, 12, 14, 20, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return n_found, n_missing


def _next_version(base_no_ext, ext):
    """ZD报价<源名><日期>_vN，同名 N+1。"""
    date = datetime.datetime.now().strftime("%Y%m%d")
    prefix = f"ZD报价{base_no_ext}{date}_v"
    v = 1
    while True:
        cand = os.path.join(C.OUTPUT_DIR, f"{prefix}{v}{ext}")
        if not os.path.exists(cand):
            return cand, v
        v += 1


def write_quote(quote_rows, system_label, ratio, source_name, progress=None):
    """生成报价单（支持多页）。quote_rows 来自 read_workbook(报价表)。

    返回 (path, version, n_found, n_missing, page_count)。
    每个源表页对应一张输出 sheet；未找到型号的行单价留空；
    每张 sheet 独立合价公式与总价 SUM。
    progress(frac, text) 在生成每张页时上报（0~1）。
    """
    cols = C.OUTPUT_COLUMNS  # list of (field, label)
    labels = [l for _, l in cols]
    fields = [f for f, _ in cols]

    price_idx = fields.index("price")
    he_idx = fields.index("he")
    qty_idx = fields.index("qty")
    name_idx = fields.index("name")

    base = os.path.splitext(os.path.basename(source_name))[0]
    out_path, version = _next_version(base, ".xlsx")

    # 按源表页分组，保持原始页顺序
    groups = {}
    for r in quote_rows:
        groups.setdefault(r.get("_sheet", "报价单"), []).append(r)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # 移除默认空 sheet
    used_titles = []
    total_found = 0
    total_missing = 0
    sheets = [s for s in groups if groups[s]]
    total_pages = len(sheets)
    page_count = 0
    for i, sheet_name in enumerate(sheets, 1):
        rows = groups[sheet_name]
        title = _safe_sheet_title(sheet_name, used_titles)
        used_titles.append(title)
        ws = wb.create_sheet(title=title)
        nf, nm = _write_sheet(ws, rows, fields, labels, price_idx,
                               he_idx, qty_idx, name_idx, system_label, ratio)
        total_found += nf
        total_missing += nm
        page_count += 1
        if progress:
            progress(i / total_pages, f"生成报价页 {i}/{total_pages}：{title}")

    if page_count == 0:
        # 无任何可识别数据：保留一张带表头的空报价单
        ws = wb.create_sheet(title="报价单")
        for ci, lab in enumerate(labels, start=1):
            c = ws.cell(row=1, column=ci, value=lab)
            c.fill = _HDR_FILL
            c.font = _HDR_FONT
            c.border = _BORDER
        page_count = 1
        if progress:
            progress(1.0, "无识别数据，已生成空报价单")

    wb.save(out_path)
    if progress:
        progress(1.0, f"报价单已保存：{os.path.basename(out_path)}")
    return out_path, version, total_found, total_missing, page_count


if __name__ == "__main__":
    print("zd_excel module loaded OK")
