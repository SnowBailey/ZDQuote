"""ZD 报价助手 - customtkinter 原生桌面界面（清爽蓝）。"""
import os
import subprocess
import threading
import tkinter.filedialog as fd
import tkinter.messagebox as mb

import customtkinter as ctk

import zd_config as C
import zd_db as DB
import zd_excel as XL
import zd_wechat as WC

ctk.set_appearance_mode("light")
ctk.set_default_color_theme("blue")  # 清爽蓝

ACCENT = "#2E5AAC"
SOFT = "#DCE6F7"
GREY = "#5A6B85"


class Prog:
    """炫彩进度指示：把 0~1 进度映射成彩虹色带，主线程 after 更新 UI。"""

    COLORS = ["#FF6B6B", "#FF922B", "#FCC419", "#51CF66",
               "#22B8CF", "#4C6EF5", "#9775FA", "#F06595"]

    def __init__(self, app, bar, label):
        self.app = app
        self.bar = bar
        self.label = label
        self._last = 0.0

    def step(self, frac, text):
        self.app.after(0, self._apply, frac, text)

    def _apply(self, frac, text):
        f = max(0.0, min(1.0, frac))
        # 单调保底：进度条只进不退，避免任何乱序回调导致抖动
        if f < self._last and frac >= 0:
            f = self._last
        else:
            self._last = f
        try:
            self.bar.set(f)
            idx = int(f * (len(self.COLORS) - 1) + 0.5) if f > 0 else 0
            self.bar.configure(progress_color=self.COLORS[idx])
        except Exception:
            pass
        try:
            self.label.configure(text=text)
        except Exception:
            pass

    def done(self, text="✅ 完成"):
        self.app.after(0, self._apply, 1.0, text)
        # 稍后盖成确认绿，避免被上面的彩虹末端色覆盖
        self.app.after(60, lambda: _safe_set(self.bar, "#20C997"))


def _safe_set(bar, color):
    try:
        bar.configure(progress_color=color)
    except Exception:
        pass


class App(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("📊 ZD 报价助手")
        self.geometry("1040x700")
        self.minsize(900, 620)
        self.last_quote_path = None

        # 标题栏
        hdr = ctk.CTkFrame(self, corner_radius=0, fg_color=ACCENT, height=58)
        hdr.pack(fill="x")
        hdr.pack_propagate(False)
        ctk.CTkLabel(hdr, text="📊  ZD 报价助手",
                      font=ctk.CTkFont(size=20, weight="bold"),
                      text_color="white").pack(side="left", padx=18, pady=12)
        ctk.CTkLabel(hdr, text="Excel 报价 · 自动清洗 · 智能查价",
                      font=ctk.CTkFont(size=12),
                      text_color="#E8EEF8").pack(side="left", padx=8)

        # 标签页
        self.tabs = ctk.CTkTabview(self, corner_radius=10)
        self.tabs.pack(fill="both", expand=True, padx=16, pady=12)
        self.t1 = self.tabs.add("① 更新数据库")
        self.t2 = self.tabs.add("② 生成报价")
        self.t3 = self.tabs.add("③ 数据库")
        self._build_tab1()
        self._build_tab2()
        self._build_tab3()

        # 状态栏
        self.status = ctk.CTkLabel(self, text="", font=ctk.CTkFont(size=12),
                                    text_color=GREY, anchor="w")
        self.status.pack(fill="x", padx=18, pady=(0, 8))
        self._refresh_status()

    # ---------- 状态栏 ----------
    def _refresh_status(self):
        DB.init_db()
        n = DB.count()
        cd = DB.get_meta("config_updated") or "—"
        pd = DB.get_meta("price_updated") or "—"
        self.status.configure(
            text=f"🗄 数据库设备数：{n}    📄 配单表更新：{cd}    💰 价格表更新：{pd}")

    # ---------- Tab1 ----------
    def _build_tab1(self):
        self.cfg_path = ctk.StringVar()
        self.prc_path = ctk.StringVar()

        f = ctk.CTkFrame(self.t1)
        f.pack(fill="both", expand=True, padx=10, pady=10)

        ctk.CTkLabel(f, text="第一步 · 上传两份源表，清洗后更新本地数据库",
                      font=ctk.CTkFont(size=14, weight="bold"),
                      text_color=ACCENT).pack(anchor="w", padx=8, pady=(6, 12))

        self._path_row(f, "📋 配单表", self.cfg_path, "选择配单表")
        self._path_row(f, "💰 价格表", self.prc_path, "选择价格表")

        self.t1_btn = ctk.CTkButton(f, text="🧹  清洗并更新数据库",
                       fg_color=ACCENT, hover_color="#234A8F",
                       command=self._do_update)
        self.t1_btn.pack(anchor="w", padx=8, pady=10)
        self.t1_bar = ctk.CTkProgressBar(f, width=440, height=14, corner_radius=7)
        self.t1_bar.pack(anchor="w", padx=8, pady=(4, 0))
        self.t1_bar.set(0)
        self.t1_plabel = ctk.CTkLabel(f, text="就绪", font=ctk.CTkFont(size=11),
                                        text_color=GREY)
        self.t1_plabel.pack(anchor="w", padx=8, pady=(2, 8))

        self.t1_log = ctk.CTkTextbox(f, height=150, corner_radius=8)
        self.t1_log.pack(fill="both", expand=True, padx=8, pady=(4, 8))
        self.t1_log.insert("0.0", "就绪。选择两份表格后点击「清洗并更新数据库」。\n")
        self.t1_log.configure(state="disabled")

    def _path_row(self, parent, label, var, btn):
        row = ctk.CTkFrame(parent, fg_color="transparent")
        row.pack(fill="x", padx=8, pady=4)
        ctk.CTkLabel(row, text=label, width=92,
                     anchor="w", font=ctk.CTkFont(size=13)).pack(side="left")
        ent = ctk.CTkEntry(row, textvariable=var,
                            placeholder_text="未选择文件", width=420)
        ent.pack(side="left", padx=6)
        ctk.CTkButton(row, text=btn, width=120, fg_color=SOFT,
                      text_color=ACCENT, hover_color="#C8D8F0",
                      command=lambda v=var: self._pick(v)).pack(side="left")

    def _pick(self, var):
        p = fd.askopenfilename(title="选择 Excel 文件",
                               filetypes=[("Excel 文件", "*.xlsx *.xls")])
        if p:
            var.set(p)

    def _do_update(self):
        cfg, prc = self.cfg_path.get(), self.prc_path.get()
        if not cfg or not prc:
            mb.showwarning("缺少文件", "请先选择配单表和价格表。")
            return
        self.t1_btn.configure(state="disabled")
        self.t1_prog = Prog(self, self.t1_bar, self.t1_plabel)
        self.t1_prog.step(0.0, "准备中…")
        self.t1_log.configure(state="normal")
        self.t1_log.delete("0.0", "end")
        self.t1_log.insert("end", "开始处理…\n")
        self.t1_log.configure(state="disabled")
        threading.Thread(target=self._update_worker, args=(cfg, prc),
                        daemon=True).start()

    def _update_worker(self, cfg, prc):
        try:
            res = XL.update_database(cfg, prc, progress=self.t1_prog.step)
            self.after(0, self._finish_update, res, None)
        except Exception as e:
            self.after(0, self._finish_update, None, str(e))

    def _finish_update(self, res, err):
        self.t1_btn.configure(state="normal")
        self.t1_log.configure(state="normal")
        if err:
            self.t1_prog.done("❌ 处理失败")
            self.t1_log.insert("end", f"❌ 出错：{err}\n")
        else:
            self.t1_prog.done("✅ 数据库更新完成")
            self.t1_log.insert("end",
                f"✅ 数据库更新完成！\n"
                f"   · 处理设备行数：{res['rows']}\n"
                f"   · 并行线程数：{res['workers']}（已按本机 CPU 核心自适应）\n"
                f"   · 配单表更新日期：{res['config_date']}\n"
                f"   · 价格表更新日期：{res['price_date']}\n")
        self.t1_log.configure(state="disabled")
        self._refresh_status()

    # ---------- Tab2 ----------
    def _build_tab2(self):
        self.q_path = ctk.StringVar()
        self.ratio = ctk.StringVar(value="1.0")
        self.sysvar = ctk.StringVar(value="大客户价（D价）")

        f = ctk.CTkFrame(self.t2)
        f.pack(fill="both", expand=True, padx=10, pady=10)

        ctk.CTkLabel(f, text="第二步 · 上传待报价表，选价格体系与比例，生成报价单",
                      font=ctk.CTkFont(size=14, weight="bold"),
                      text_color=ACCENT).pack(anchor="w", padx=8, pady=(6, 12))

        self._path_row(f, "🧾 待报价表", self.q_path, "选择报价表")

        row = ctk.CTkFrame(f, fg_color="transparent")
        row.pack(fill="x", padx=8, pady=4)
        ctk.CTkLabel(row, text="💲 价格体系", width=92, anchor="w").pack(side="left")
        ctk.CTkComboBox(row, values=list(C.PRICE_SYSTEMS.keys()),
                         variable=self.sysvar, width=130,
                         state="readonly").pack(side="left", padx=6)
        ctk.CTkLabel(row, text="✖ 比例", width=60, anchor="w").pack(side="left", padx=(10, 0))
        ctk.CTkEntry(row, textvariable=self.ratio, width=90).pack(side="left", padx=6)
        ctk.CTkLabel(row, text="例：大客户价（D价）、比例0.9 → 单价=D价×0.9",
                      text_color=GREY, font=ctk.CTkFont(size=11)).pack(side="left", padx=8)

        self.t2_btn = ctk.CTkButton(f, text="📤  生成报价单",
                       fg_color=ACCENT, hover_color="#234A8F",
                       command=self._do_quote)
        self.t2_btn.pack(anchor="w", padx=8, pady=10)

        act = ctk.CTkFrame(f, fg_color="transparent")
        act.pack(fill="x", padx=8, pady=2)
        ctk.CTkButton(act, text="📂 打开文件", width=130, fg_color=SOFT,
                       text_color=ACCENT, hover_color="#C8D8F0",
                       command=self._open_last).pack(side="left", padx=4)
        ctk.CTkButton(act, text="📁 打开输出文件夹", width=150, fg_color=SOFT,
                       text_color=ACCENT, hover_color="#C8D8F0",
                       command=self._open_outdir).pack(side="left", padx=4)
        ctk.CTkButton(act, text="💬 转发至微信", width=130, fg_color="#07C160",
                       text_color="white", hover_color="#06AD56",
                       command=self._forward).pack(side="left", padx=4)

        self.t2_bar = ctk.CTkProgressBar(f, width=440, height=14, corner_radius=7)
        self.t2_bar.pack(anchor="w", padx=8, pady=(4, 0))
        self.t2_bar.set(0)
        self.t2_plabel = ctk.CTkLabel(f, text="就绪", font=ctk.CTkFont(size=11),
                                         text_color=GREY)
        self.t2_plabel.pack(anchor="w", padx=8, pady=(2, 8))

        self.t2_log = ctk.CTkTextbox(f, height=150, corner_radius=8)
        self.t2_log.pack(fill="both", expand=True, padx=8, pady=(4, 8))
        self.t2_log.insert("0.0", "就绪。\n")
        self.t2_log.configure(state="disabled")

    def _do_quote(self):
        qp = self.q_path.get()
        if not qp:
            mb.showwarning("缺少文件", "请先选择待报价表。")
            return
        try:
            ratio = float(self.ratio.get())
        except ValueError:
            mb.showwarning("比例错误", "比例必须是数字，例如 0.9。")
            return
        sys = self.sysvar.get()
        self.t2_btn.configure(state="disabled")
        self.t2_prog = Prog(self, self.t2_bar, self.t2_plabel)
        self.t2_prog.step(0.0, "准备中…")
        self.t2_log.configure(state="normal")
        self.t2_log.delete("0.0", "end")
        self.t2_log.insert("end", f"开始处理：{os.path.basename(qp)}\n")
        self.t2_log.configure(state="disabled")
        threading.Thread(target=self._quote_worker, args=(qp, sys, ratio),
                        daemon=True).start()

    def _quote_worker(self, qp, sys, ratio):
        try:
            self.t2_prog.step(0.02, f"读取待报价表：{os.path.basename(qp)}")
            rows = XL.read_workbook(
                qp, progress=XL.scaled(self.t2_prog.step, 0.02, 0.5))
            out, ver, nf, nm, pages = XL.write_quote(
                rows, sys, ratio, qp,
                progress=XL.scaled(self.t2_prog.step, 0.5, 1.0))
            self.after(0, self._finish_quote, out, ver, nf, nm, pages, None)
        except Exception as e:
            self.after(0, self._finish_quote, None, 0, 0, 0, 0, str(e))

    def _finish_quote(self, out, ver, nf, nm, pages, err):
        self.t2_btn.configure(state="normal")
        self.t2_log.configure(state="normal")
        if err:
            self.t2_prog.done("❌ 处理失败")
            self.t2_log.insert("end", f"❌ 出错：{err}\n")
        else:
            self.last_quote_path = out
            self.t2_prog.done("✅ 报价单已生成")
            self.t2_log.insert("end",
                f"✅ 报价单已生成（第 {ver} 版）\n"
                f"   · 路径：{out}\n"
                f"   · 生成分表页：{pages} 页（对应源表各页）\n"
                f"   · 命中单价：{nf} 行，未找到型号：{nm} 行（单价留空）\n")
        self.t2_log.configure(state="disabled")
        self._refresh_status()

    def _open_last(self):
        if self.last_quote_path and os.path.exists(self.last_quote_path):
            subprocess.run(["open", self.last_quote_path], check=False)
        else:
            mb.showinfo("提示", "尚未生成报价单。")

    def _open_outdir(self):
        subprocess.run(["open", C.OUTPUT_DIR], check=False)

    def _forward(self):
        if not self.last_quote_path or not os.path.exists(self.last_quote_path):
            mb.showinfo("提示", "请先生成报价单再转发。")
            return
        msg = WC.forward(self.last_quote_path)
        mb.showinfo("转发微信", msg)

    # ---------- Tab3 ----------
    def _build_tab3(self):
        f = ctk.CTkFrame(self.t3)
        f.pack(fill="both", expand=True, padx=10, pady=10)
        ctk.CTkLabel(f, text="数据库状态与管理",
                      font=ctk.CTkFont(size=14, weight="bold"),
                      text_color=ACCENT).pack(anchor="w", padx=8, pady=(6, 12))

        self.db_info = ctk.CTkLabel(f, text="", justify="left",
                                     font=ctk.CTkFont(size=13))
        self.db_info.pack(anchor="w", padx=12, pady=6)

        btns = ctk.CTkFrame(f, fg_color="transparent")
        btns.pack(anchor="w", padx=8, pady=10)
        ctk.CTkButton(btns, text="🔄 刷新状态", width=130, fg_color=SOFT,
                       text_color=ACCENT, hover_color="#C8D8F0",
                       command=self._refresh_db_info).pack(side="left", padx=4)
        ctk.CTkButton(btns, text="📤 导出数据库为 Excel", width=180,
                       fg_color=SOFT, text_color=ACCENT,
                       hover_color="#C8D8F0",
                       command=self._export_db).pack(side="left", padx=4)
        ctk.CTkButton(btns, text="🗑 清空数据库", width=130,
                       fg_color="#E5534B", text_color="white",
                       hover_color="#C73E36",
                       command=self._clear_db).pack(side="left", padx=4)
        self._refresh_db_info()

    def _refresh_db_info(self):
        DB.init_db()
        n = DB.count()
        cd = DB.get_meta("config_updated") or "—"
        pd = DB.get_meta("price_updated") or "—"
        self.db_info.configure(
            text=f"设备总数：{n}\n配单表更新日期：{cd}\n价格表更新日期：{pd}")

    def _export_db(self):
        DB.init_db()
        rows = DB.all_rows()
        if not rows:
            mb.showinfo("提示", "数据库为空。")
            return
        out = os.path.join(C.OUTPUT_DIR, "数据库导出.xlsx")
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        wb = openpyxl.Workbook()
        ws = wb.active
        cols = ["序号", "设备名称", "品牌", "型号", "规格", "国别", "数量",
                 "单位", "备注", "产品编码", "标底参数",
                 "核心价(H)", "签约价(Q)", "大客户价(D)", "批发价(P)", "业务底价(Y)",
                 "概算价", "销售指导价", "租赁价", "市场价"]
        for ci, c in enumerate(cols, 1):
            cell = ws.cell(row=1, column=ci, value=c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2E5AAC")
        keymap = ["seq", "name", "brand", "model", "spec", "country", "qty",
                   "unit", "remark", "code", "bid_param",
                   "h_price", "q_price", "d_price", "p_price", "y_price",
                   "budget_price", "sale_guide_price", "lease_price", "market_price"]
        for ri, r in enumerate(rows, 2):
            for ci, k in enumerate(keymap, 1):
                ws.cell(row=ri, column=ci, value=r.get(k))
        wb.save(out)
        mb.showinfo("导出完成", f"已导出到：{out}")
        subprocess.run(["open", out], check=False)

    def _clear_db(self):
        if not mb.askyesno("确认", "确定清空数据库？此操作不可恢复。"):
            return
        DB.clear_db()
        self._refresh_db_info()
        self._refresh_status()
        mb.showinfo("完成", "数据库已清空。")


def main():
    DB.init_db()
    app = App()
    app.mainloop()


if __name__ == "__main__":
    main()
